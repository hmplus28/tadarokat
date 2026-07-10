import tempfile
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
from django.contrib.auth import get_user_model
from django.test import TestCase

from services.excel_importer import ExcelImporter
from services.excel_template_service import ExcelTemplateService

User = get_user_model()


class ExcelTemplateServiceTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='tpl_admin',
            email='tpl@test.local',
            password='pass',
            role='admin',
        )

    def test_template_headers_match_importer_schema(self):
        specs = ExcelTemplateService.get_column_specs()
        import_cols = ExcelImporter.get_import_columns()
        self.assertEqual(len(specs), len(import_cols))
        self.assertEqual(
            [s['header'] for s in specs],
            [h for h, _ in import_cols],
        )
        self.assertIn('شماره درخواست خرید', [s['header'] for s in specs])

    def test_generated_workbook_has_guide_and_meta_sheets(self):
        content, filename, version = ExcelTemplateService.generate_template_bytes()
        self.assertTrue(filename.endswith('.xlsx'))
        self.assertEqual(len(version), 12)

        wb = openpyxl.load_workbook(BytesIO(content))
        self.assertEqual(wb.sheetnames, ['ورودی', 'راهنما', 'نسخه قالب'])
        ws = wb['ورودی']
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers[0], 'شماره درخواست خرید')
        self.assertEqual(ws['A2'].value, '10001')

    def test_template_importable_after_removing_sample_row(self):
        content, _, _ = ExcelTemplateService.generate_template_bytes()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'template.xlsx'
            path.write_bytes(content)

            wb = openpyxl.load_workbook(path)
            ws = wb.active
            ws.delete_rows(2)
            wb.save(path)

            stats = ExcelImporter.import_from_excel(str(path), user=self.user)
            self.assertEqual(stats['created'], 0)
            self.assertEqual(stats['updated'], 0)

    def test_schema_version_changes_when_column_map_changes(self):
        v1 = ExcelTemplateService.get_schema_version()
        original = dict(ExcelImporter.COLUMN_MAP)
        try:
            ExcelImporter.COLUMN_MAP['ستون تست'] = 'description'
            v2 = ExcelTemplateService.get_schema_version()
            self.assertNotEqual(v1, v2)
        finally:
            ExcelImporter.COLUMN_MAP.clear()
            ExcelImporter.COLUMN_MAP.update(original)

    def test_template_readable_by_pandas(self):
        content, _, _ = ExcelTemplateService.generate_template_bytes()
        df = pd.read_excel(BytesIO(content), sheet_name=0, engine='openpyxl')
        df = ExcelImporter.normalize_columns(df)
        self.assertIn('شماره درخواست خرید', df.columns)