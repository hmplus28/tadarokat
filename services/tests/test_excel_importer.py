"""
سناریوهای تست آپلود و همگام‌سازی اکسل

دو فایل اکسل:
1. incoming — اکسل ورودی کاربر (آپلود)
2. panel/output — همان فایل پس از غنی‌سازی با داده‌های پنل (export_to_excel)

دیتابیس مرجع است؛ import حذف نمی‌کند، فقط UPSERT می‌کند.
"""
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
from django.contrib.auth import get_user_model
from django.test import TestCase

from accounts.models import ProductCategory
from deliveries.models import Delivery
from inquiries.models import Inquiry
from orders.models import Order
from purchases.models import Purchase
from services.excel_importer import ExcelImporter, ExcelImportError

User = get_user_model()

EXCEL_HEADERS = [
    'شماره درخواست خرید',
    'نام قلم خریدنی',
    'کد قلم خریدنی',
    'مقدار درخواست',
    'واحد سنجش',
    'شماره استعلام',
    'شماره دستور خرید',
    'تاریخ دستور خرید',
    'شماره تحویل',
    'تاریخ تحویل',
]


def _write_excel(path: Path, rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(EXCEL_HEADERS)
    for row in rows:
        ws.append([
            row.get('purchase_number', ''),
            row.get('product_title', ''),
            row.get('product_code', ''),
            row.get('quantity', ''),
            row.get('unit', ''),
            row.get('inquiry_number', ''),
            row.get('order_number', ''),
            row.get('order_date', ''),
            row.get('delivery_number', ''),
            row.get('delivery_date', ''),
        ])
    wb.save(path)


class ExcelImporterScenarioTests(TestCase):
    """تست سناریوهای آپلود، همگام‌سازی و حفظ داده"""

    def setUp(self):
        self.user = User.objects.create_user(
            username='excel_admin',
            email='excel@test.com',
            password='pass',
            role='admin',
        )
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.excel_path = Path(self.tmp_dir.name) / 'test_input.xlsx'

    def tearDown(self):
        self.tmp_dir.cleanup()

    def _import(self, rows, full_sync=False):
        _write_excel(self.excel_path, rows)
        return ExcelImporter.import_from_excel(
            str(self.excel_path),
            user=self.user,
            full_sync=full_sync,
        )

    def test_scenario_1_first_upload_creates_purchases(self):
        """آپلود اول — رکوردهای جدید در DB ساخته می‌شوند"""
        stats = self._import([
            {'purchase_number': '9001', 'product_title': 'کالای الف', 'quantity': 2, 'unit': 'عدد'},
            {'purchase_number': '9002', 'product_title': 'کالای ب', 'quantity': 3, 'unit': 'عدد'},
        ])
        self.assertEqual(stats['created'], 2)
        self.assertEqual(stats['updated'], 0)
        self.assertEqual(Purchase.objects.filter(purchase_number__in=['9001', '9002']).count(), 2)

    def test_scenario_2_reupload_updates_non_workflow_fields(self):
        """همگام‌سازی مجدد — فیلدهای غیرگردش‌کار از اکسل به‌روز می‌شوند"""
        self._import([
            {'purchase_number': '9010', 'product_title': 'نسخه اول', 'quantity': 1, 'unit': 'عدد'},
        ])
        stats = self._import([
            {'purchase_number': '9010', 'product_title': 'نسخه دوم', 'quantity': 5, 'unit': 'کیلو'},
        ])
        purchase = Purchase.objects.get(purchase_number='9010')
        self.assertEqual(stats['updated'], 1)
        self.assertEqual(purchase.product_title, 'نسخه دوم')
        self.assertEqual(float(purchase.quantity), 5.0)
        self.assertEqual(purchase.unit, 'کیلو')

    def test_scenario_3_panel_workflow_preserved_on_smart_sync(self):
        """داده ثبت‌شده در پنل — فیلدهای workflow روی اکسل اولویت دارند"""
        self._import([
            {
                'purchase_number': '9020',
                'product_title': 'کالا',
                'inquiry_number': 'INQ-OLD',
                'order_number': 'ORD-OLD',
            },
        ])
        purchase = Purchase.objects.get(purchase_number='9020')
        purchase.inquiry_number = 'INQ-PANEL'
        purchase.order_number = 'ORD-PANEL'
        purchase.order_date = '1405/04/01'
        purchase.save()

        stats = self._import([
            {
                'purchase_number': '9020',
                'product_title': 'کالا',
                'inquiry_number': 'INQ-EXCEL-NEW',
                'order_number': 'ORD-EXCEL-NEW',
                'order_date': '1405/01/01',
            },
        ])
        purchase.refresh_from_db()
        self.assertGreater(stats['workflow_preserved'], 0)
        self.assertEqual(purchase.inquiry_number, 'INQ-PANEL')
        self.assertEqual(purchase.order_number, 'ORD-PANEL')
        self.assertEqual(purchase.order_date, '1405/04/01')

    def test_scenario_4_purchases_not_in_excel_are_not_deleted(self):
        """اکسل کوچک‌تر — پرونده‌های پنل حذف نمی‌شوند"""
        self._import([
            {'purchase_number': '9031', 'product_title': 'در اکسل'},
            {'purchase_number': '9032', 'product_title': 'در اکسل ۲'},
        ])
        Purchase.objects.create(
            purchase_number='9033',
            line_number=1,
            product_title='فقط در پنل',
        )
        self.assertEqual(Purchase.objects.count(), 3)

        stats = self._import([
            {'purchase_number': '9031', 'product_title': 'در اکسل'},
        ])
        self.assertEqual(Purchase.objects.count(), 3)
        self.assertTrue(Purchase.objects.filter(purchase_number='9033').exists())
        self.assertEqual(stats['updated'], 1)

    def test_scenario_5_panel_only_purchase_appended_to_excel(self):
        """پرونده فقط پنل — به اکسل خروجی اضافه می‌شود"""
        self._import([
            {'purchase_number': '9041', 'product_title': 'ردیف اکسل'},
        ])
        Purchase.objects.create(
            purchase_number='9042',
            line_number=1,
            product_title='ردیف پنل',
            product_code='P-99',
        )
        export_stats = ExcelImporter.export_to_excel(str(self.excel_path))
        self.assertGreaterEqual(export_stats['rows_appended'], 1)

        df = pd.read_excel(self.excel_path, engine='openpyxl')
        df = ExcelImporter.normalize_columns(df)
        df = ExcelImporter.map_columns(df)
        titles = df['product_title'].astype(str).tolist()
        self.assertIn('ردیف پنل', titles)

    def test_scenario_6_inquiry_sync_no_duplicate(self):
        """استعلام — دوباره ساخته نمی‌شود"""
        self._import([
            {'purchase_number': '9051', 'product_title': 'کالا', 'inquiry_number': 'INQ-9051'},
        ])
        first = Inquiry.objects.filter(inquiry_number='INQ-9051').count()
        stats = self._import([
            {'purchase_number': '9051', 'product_title': 'کالا', 'inquiry_number': 'INQ-9051'},
        ])
        second = Inquiry.objects.filter(inquiry_number='INQ-9051').count()
        self.assertEqual(first, 1)
        self.assertEqual(second, 1)
        self.assertGreater(stats['inquiries_skipped'], 0)

    def test_scenario_7_order_sync_no_duplicate(self):
        """دستور خرید — دوباره ساخته نمی‌شود"""
        self._import([
            {
                'purchase_number': '9061',
                'product_title': 'کالا',
                'inquiry_number': 'INQ-9061',
                'order_number': 'ORD-9061',
                'order_date': '1405/02/01',
            },
        ])
        first = Order.objects.filter(order_number='ORD-9061').count()
        stats = self._import([
            {
                'purchase_number': '9061',
                'product_title': 'کالا',
                'inquiry_number': 'INQ-9061',
                'order_number': 'ORD-9061',
                'order_date': '1405/02/01',
            },
        ])
        second = Order.objects.filter(order_number='ORD-9061').count()
        self.assertEqual(first, 1)
        self.assertEqual(second, 1)
        self.assertGreater(stats['orders_skipped'], 0)

    def test_scenario_8_delivery_sync_no_duplicate(self):
        """تحویل — دوباره ساخته نمی‌شود"""
        self._import([
            {
                'purchase_number': '9071',
                'product_title': 'کالا',
                'order_number': 'ORD-9071',
                'delivery_number': 'DLV-9071',
                'delivery_date': '1405/03/01',
            },
        ])
        first = Delivery.objects.filter(delivery_number='DLV-9071').count()
        stats = self._import([
            {
                'purchase_number': '9071',
                'product_title': 'کالا',
                'order_number': 'ORD-9071',
                'delivery_number': 'DLV-9071',
                'delivery_date': '1405/03/01',
            },
        ])
        second = Delivery.objects.filter(delivery_number='DLV-9071').count()
        self.assertEqual(first, 1)
        self.assertEqual(second, 1)
        self.assertGreater(stats['deliveries_skipped'], 0)

    def test_scenario_9_export_enriches_excel_with_panel_columns(self):
        """اکسل خروجی — ستون‌های پنل (انحراف، وضعیت) اضافه می‌شوند"""
        self._import([
            {'purchase_number': '9081', 'product_title': 'کالا تست'},
        ])
        ExcelImporter.export_to_excel(str(self.excel_path))
        wb = openpyxl.load_workbook(self.excel_path)
        headers = [
            ExcelImporter._normalize_header(c.value)
            for c in wb.active[1]
        ]
        self.assertIn('انحراف مهلت', headers)
        self.assertIn('وضعیت فعلی پنل', headers)
        self.assertIn('گروه بندی کالایی', headers)

    def test_scenario_11_import_creates_new_product_category(self):
        """گروه کالایی جدید در اکسل — در سیستم ثبت می‌شود"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['شماره درخواست خرید', 'نام قلم خریدنی', 'گروه بندی کالایی'])
        ws.append(['9101', 'کالای جدید', 'مواد شیمیایی تست'])
        wb.save(self.excel_path)

        stats = ExcelImporter.import_from_excel(str(self.excel_path), user=self.user)
        self.assertEqual(stats.get('categories_created'), 1)
        self.assertTrue(ProductCategory.objects.filter(name='مواد شیمیایی تست').exists())
        purchase = Purchase.objects.get(purchase_number='9101')
        self.assertEqual(purchase.product_category, 'مواد شیمیایی تست')

    def test_scenario_12_export_writes_product_category(self):
        """گروه کالایی پرونده — در اکسل خروجی نوشته می‌شود"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['شماره درخواست خرید', 'نام قلم خریدنی'])
        ws.append(['9111', 'کالا با گروه'])
        wb.save(self.excel_path)

        ExcelImporter.import_from_excel(str(self.excel_path), user=self.user)
        purchase = Purchase.objects.get(purchase_number='9111')
        purchase.product_category = 'آی تی'
        purchase.save(update_fields=['product_category'])

        ExcelImporter.export_to_excel(str(self.excel_path))
        wb = openpyxl.load_workbook(self.excel_path)
        headers = [ExcelImporter._normalize_header(c.value) for c in wb.active[1]]
        self.assertIn('گروه بندی کالایی', headers)

        category_col = headers.index('گروه بندی کالایی') + 1
        value = wb.active.cell(row=2, column=category_col).value
        self.assertEqual(str(value).strip(), 'آی تی')

    def test_scenario_13_atomic_rollback_on_fail_on_errors(self):
        """fail_on_errors=True — خطا باعث rollback کامل می‌شود"""
        from unittest.mock import patch

        _write_excel(self.excel_path, [
            {'purchase_number': '9201', 'product_title': 'ردیف اول', 'quantity': 1, 'unit': 'عدد'},
            {'purchase_number': '9202', 'product_title': 'ردیف دوم', 'quantity': 2, 'unit': 'عدد'},
        ])

        original_create = Purchase.objects.create

        def flaky_create(*args, **kwargs):
            if kwargs.get('purchase_number') == '9202':
                raise ValueError('simulated row error')
            return original_create(*args, **kwargs)

        before_count = Purchase.objects.count()
        with patch.object(Purchase.objects, 'create', side_effect=flaky_create):
            with self.assertRaises(ExcelImportError):
                ExcelImporter.import_from_excel(
                    str(self.excel_path),
                    user=self.user,
                    fail_on_errors=True,
                )

        self.assertEqual(Purchase.objects.count(), before_count)
        self.assertFalse(Purchase.objects.filter(purchase_number__in=['9201', '9202']).exists())

    def test_scenario_14_export_uses_separate_output_path(self):
        """export_path جدا — خروجی روی مسیر panel نوشته می‌شود (ورودی پس از پردازش آزاد می‌شود)"""
        incoming = Path(self.tmp_dir.name) / 'incoming.xlsx'
        panel = Path(self.tmp_dir.name) / 'panel.xlsx'
        _write_excel(incoming, [
            {'purchase_number': '9211', 'product_title': 'تست خروجی', 'quantity': 2, 'unit': 'عدد'},
        ])
        ExcelImporter.import_from_excel(
            str(incoming),
            user=self.user,
            export_path=str(panel),
        )
        self.assertTrue(panel.exists())
        self.assertFalse(incoming.exists())
        wb = openpyxl.load_workbook(panel)
        headers = [ExcelImporter._normalize_header(c.value) for c in wb.active[1]]
        self.assertIn('وضعیت فعلی پنل', headers)

    def test_scenario_10_full_sync_overwrites_workflow_from_excel(self):
        """full_sync=True — فیلدهای workflow از اکسل بازنویسی می‌شوند (حالت خطرناک)"""
        self._import([
            {'purchase_number': '9091', 'product_title': 'کالا', 'order_number': 'ORD-A'},
        ])
        purchase = Purchase.objects.get(purchase_number='9091')
        purchase.order_number = 'ORD-PANEL'
        purchase.save()

        self._import(
            [{'purchase_number': '9091', 'product_title': 'کالا', 'order_number': 'ORD-EXCEL'}],
            full_sync=True,
        )
        purchase.refresh_from_db()
        self.assertEqual(purchase.order_number, 'ORD-EXCEL')