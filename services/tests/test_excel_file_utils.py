import tempfile
from pathlib import Path

import openpyxl
from django.test import TestCase

from services.excel_file_utils import (
    load_workbook_from_bytes,
    read_file_bytes,
    resolve_panel_export_path,
    safe_copy_bytes,
    safe_save_workbook,
    workbook_to_bytes,
)


class ExcelFileUtilsTests(TestCase):
    def test_roundtrip_bytes_load_and_safe_save(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / 'source.xlsx'
            dst = Path(tmp) / 'output.xlsx'

            wb = openpyxl.Workbook()
            wb.active.append(['شماره درخواست خرید', 'نام قلم خریدنی'])
            wb.active.append(['1001', 'کالا'])
            wb.save(src)
            wb.close()

            data = read_file_bytes(src)
            loaded = load_workbook_from_bytes(data)
            loaded.active.cell(row=2, column=2, value='کالای ویرایش‌شده')
            written = safe_save_workbook(loaded, dst)
            self.assertEqual(written, dst)

            result = openpyxl.load_workbook(dst)
            self.assertEqual(result.active.cell(row=2, column=2).value, 'کالای ویرایش‌شده')
            result.close()

    def test_safe_copy_bytes(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / 'a.xlsx'
            dst = Path(tmp) / 'b.xlsx'
            wb = openpyxl.Workbook()
            wb.save(src)
            wb.close()

            safe_copy_bytes(read_file_bytes(src), dst)
            self.assertTrue(dst.exists())
            self.assertGreater(dst.stat().st_size, 0)

    def test_workbook_to_bytes_never_touches_target_dir(self):
        wb = openpyxl.Workbook()
        wb.active.append(['test'])
        data = workbook_to_bytes(wb)
        self.assertGreater(len(data), 100)

    def test_resolve_panel_export_path_prefers_root_panel(self):
        from django.conf import settings
        from services.excel_file_utils import get_panel_excel_path

        panel_root = get_panel_excel_path()
        self.assertEqual(resolve_panel_export_path('any_token'), panel_root)

    def test_resolve_panel_export_path_finds_suffixed_file(self):
        from django.conf import settings
        from services.excel_file_utils import get_panel_excel_path

        panel_root = get_panel_excel_path()
        panel_backup = panel_root.read_bytes() if panel_root.exists() else None
        if panel_root.exists():
            panel_root.unlink()

        upload_dir = Path(settings.MEDIA_ROOT) / 'excel_uploads'
        upload_dir.mkdir(parents=True, exist_ok=True)
        token = '20990101_120000'
        path = upload_dir / f'panel_{token}_abc12345.xlsx'
        wb = openpyxl.Workbook()
        wb.save(path)
        wb.close()
        try:
            resolved = resolve_panel_export_path(token)
            self.assertEqual(resolved, path)
        finally:
            path.unlink(missing_ok=True)
            if panel_backup is not None:
                panel_root.write_bytes(panel_backup)