from decimal import Decimal
from io import BytesIO

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase

from purchases.models import Purchase
from services.excel_full_export_service import ExcelFullExportService
from services.excel_importer import ExcelImporter

User = get_user_model()


class ExcelFullExportServiceTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='export_admin',
            email='export@test.local',
            password='pass',
            role='admin',
        )
        Purchase.objects.create(
            purchase_number='8001',
            line_number=1,
            product_title='کالای تست',
            product_category='آی تی',
            quantity=Decimal('5'),
            unit='عدد',
        )
        Purchase.objects.create(
            purchase_number='8002',
            line_number=1,
            product_title='کالای دوم',
            quantity=Decimal('1'),
        )
        Purchase.objects.filter(purchase_number='8001').update(
            deadline_deviated=True,
            order_deadline_deviated=False,
            payment_lead_deviated=True,
            current_status=Purchase.CurrentStatus.INQUIRY_ISSUED,
        )

    def test_export_columns_include_panel_fields(self):
        import_cols = {f for _, f in ExcelImporter.get_import_columns()}
        export_cols = ExcelImporter.get_export_columns()
        export_fields = [f for _, f in export_cols]
        self.assertGreater(len(export_cols), len(import_cols))
        self.assertIn('_current_status_label', export_fields)
        self.assertIn('_payment_lead_deviated_label', export_fields)
        self.assertEqual(export_cols[0][0], 'شماره درخواست خرید')

    def test_full_export_workbook_structure(self):
        content, filename, row_count = ExcelFullExportService.generate_export_bytes()
        self.assertTrue(filename.endswith('.xlsx'))
        self.assertEqual(row_count, 2)

        wb = openpyxl.load_workbook(BytesIO(content))
        self.assertEqual(wb.sheetnames, ['خروجی', 'راهنما', 'نسخه قالب'])

        ws = wb['خروجی']
        headers = [cell.value for cell in ws[1]]
        expected_headers = [h for h, _ in ExcelImporter.get_export_columns()]
        self.assertEqual(headers, expected_headers)
        self.assertIn('وضعیت فعلی پنل', headers)
        self.assertIn('انحراف لید تایم پرداخت', headers)

        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 2)

        purchase_col = headers.index('شماره درخواست خرید')
        numbers = {row[purchase_col] for row in data_rows}
        self.assertEqual(numbers, {'8001', '8002'})

    def test_panel_computed_fields_in_export(self):
        content, _, _ = ExcelFullExportService.generate_export_bytes()
        wb = openpyxl.load_workbook(BytesIO(content))
        ws = wb['خروجی']
        headers = [cell.value for cell in ws[1]]

        status_col = headers.index('وضعیت فعلی پنل')
        deadline_col = headers.index('انحراف مهلت')
        payment_col = headers.index('انحراف لید تایم پرداخت')

        row_8001 = next(
            row for row in ws.iter_rows(min_row=2, values_only=True)
            if row[headers.index('شماره درخواست خرید')] == '8001'
        )
        self.assertEqual(row_8001[status_col], 'استعلام صادر شده')
        self.assertEqual(row_8001[deadline_col], 'بله')
        self.assertEqual(row_8001[payment_col], 'بله')