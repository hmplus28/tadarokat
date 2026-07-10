"""تولید قالب اکسل ورودی — داینامیک بر اساس ExcelImporter.COLUMN_MAP"""
from __future__ import annotations

import hashlib
from datetime import datetime
from io import BytesIO
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.excel_importer import ExcelImporter


FIELD_HINTS: Dict[str, str] = {
    'purchase_number': 'الزامی — یکتا برای هر درخواست؛ چند خط با همان شماره = چند قلم',
    'base_number': 'شماره درخواست کالا از انبار',
    'request_date': 'تاریخ شمسی — مثال: 1405/04/01',
    'purchase_date': 'تاریخ ثبت درخواست خرید (ستون D در گزارش‌ها)',
    'supply_unit': 'انبار / واحد تامین مقصد',
    'product_category': 'گروه کالایی — مثال: آی تی، ابنیه و مصالح',
    'requester': 'نام درخواست‌کننده',
    'product_code': 'کد قلم خریدنی',
    'product_title': 'عنوان کالا',
    'unit': 'واحد سنجش — عدد، کیلوگرم، ...',
    'quantity': 'مقدار عددی',
    'expert_name': 'کارشناس مسئول',
    'description': 'توضیحات',
    'status': 'وضعیت درخواست — مثال: در جريان',
    'required_date': 'تاریخ نیاز',
    'inquiry_deadline': 'مهلت استعلام (ستون O)',
    'inquiry_deadline_2': 'مهلت استعلام دوم',
    'inquiry_number': 'شماره استعلام صادرشده',
    'inquiry_received_date': 'تاریخ دریافت درخواست',
    'purchase_type': 'نوع روند خرید',
    'preinvoice_number': 'شماره پیش‌فاکتور',
    'preinvoice_price': 'فی پیش‌فاکتور (ریال)',
    'preinvoice_total': 'جمع پیش‌فاکتور',
    'order_number': 'شماره دستور خرید',
    'order_date': 'تاریخ دستور (ستون AB)',
    'order_request_number': 'شماره سفارش',
    'order_request_status': 'وضعیت سفارش',
    'order_request_date': 'تاریخ سفارش',
    'order_total': 'جمع کل سفارش (ستون AC)',
    'advance_payment': 'پیش‌پرداخت (ستون AD)',
    'deductions': 'کسور / تخفیف پیش‌فاکتور (ریال)',
    'delivered_quantity': 'مقدار تحویل‌شده',
    'delivery_number': 'شماره تحویل / رسید',
    'delivery_date': 'تاریخ تحویل',
    'supplier': 'تامین‌کننده',
    'invoice_price': 'فی فاکتور',
    'invoice_number': 'شماره فاکتور',
    'invoice_total': 'جمع فاکتور (ستون AL)',
    'payment_request_amount': 'مبلغ درخواست پرداخت (ستون AN)',
    'payment_request_number': 'شماره درخواست پرداخت',
    'payment_registration_date': 'تاریخ ثبت واریزی (AP)',
    'payment_completion_date': 'تاریخ انجام واریزی (AQ)',
}

SAMPLE_VALUES: Dict[str, object] = {
    'purchase_number': '10001',
    'base_number': 'WR-1405-001',
    'request_date': '1405/04/01',
    'purchase_date': '1405/04/02',
    'supply_unit': 'انبار مصرفی',
    'product_category': 'آی تی',
    'requester': 'نام درخواست‌کننده',
    'product_code': '410001',
    'product_title': 'نمونه کالا — این سطر را حذف یا ویرایش کنید',
    'unit': 'عدد',
    'quantity': 2,
    'expert_name': 'کارشناس نمونه',
    'description': 'توضیح نمونه',
    'status': 'در جريان',
    'required_date': '1405/05/01',
}


class ExcelTemplateService:
    """ساخت فایل قالب ورودی از اسکیمای زنده importer"""

    @classmethod
    def get_schema_version(cls) -> str:
        """هش اسکیمای زنده — با هر تغییر در COLUMN_MAP / alias / الزامی‌ها عوض می‌شود."""
        cols = ExcelImporter.get_import_columns()
        full_map = sorted(f'{h}:{f}' for h, f in ExcelImporter.COLUMN_MAP.items())
        aliases = sorted(f'{h}:{f}' for h, f in ExcelImporter.COLUMN_ALIASES.items())
        required = sorted(ExcelImporter.IMPORT_REQUIRED_FIELDS)
        payload = '\n'.join([
            'cols:' + '|'.join(f'{h}:{f}' for h, f in cols),
            'map:' + '|'.join(full_map),
            'alias:' + '|'.join(aliases),
            'req:' + '|'.join(required),
        ])
        return hashlib.sha256(payload.encode('utf-8')).hexdigest()[:12]

    @classmethod
    def get_column_specs(cls) -> List[dict]:
        specs = []
        for header, field in ExcelImporter.get_import_columns():
            specs.append({
                'header': header,
                'field': field,
                'required': field in ExcelImporter.IMPORT_REQUIRED_FIELDS,
                'decimal': field in ExcelImporter.DECIMAL_FIELDS,
                'workflow': field in ExcelImporter.WORKFLOW_FIELDS,
                'aliases': ExcelImporter.get_import_aliases_for_field(field),
                'hint': FIELD_HINTS.get(field, ''),
            })
        return specs

    @classmethod
    def build_workbook(cls) -> Tuple[Workbook, str]:
        wb = Workbook()
        ws = wb.active
        ws.title = 'ورودی'

        specs = cls.get_column_specs()
        headers = [s['header'] for s in specs]

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill('solid', fgColor='4F46E5')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        ws.append(headers)
        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        sample_row = []
        for spec in specs:
            val = SAMPLE_VALUES.get(spec['field'], '')
            sample_row.append(val)
        ws.append(sample_row)

        for col_idx, header in enumerate(headers, start=1):
            width = max(len(header) + 2, 14)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 36)

        ws.freeze_panes = 'A2'

        cls._write_guide_sheet(wb, specs)
        cls._write_meta_sheet(wb, specs)

        version = cls.get_schema_version()
        return wb, version

    @classmethod
    def _write_guide_sheet(cls, wb: Workbook, specs: List[dict]) -> None:
        ws = wb.create_sheet('راهنما')
        ws.append(['ستون اکسل', 'فیلد سیستم', 'الزامی', 'نوع', 'نام‌های جایگزین', 'راهنما'])
        for spec in specs:
            type_label = 'عدد' if spec['decimal'] else 'متن/تاریخ'
            if spec['workflow']:
                type_label += ' · گردش‌کار'
            aliases = '، '.join(spec['aliases']) if spec['aliases'] else '—'
            ws.append([
                spec['header'],
                spec['field'],
                'بله' if spec['required'] else 'خیر',
                type_label,
                aliases,
                spec['hint'] or '—',
            ])
        for col in range(1, 7):
            ws.column_dimensions[get_column_letter(col)].width = 22 if col < 6 else 40
        ws.freeze_panes = 'A2'

    @classmethod
    def _write_meta_sheet(cls, wb: Workbook, specs: List[dict]) -> None:
        ws = wb.create_sheet('نسخه قالب')
        version = cls.get_schema_version()
        ws.append(['کلید', 'مقدار'])
        ws.append(['نسخه اسکیما', version])
        ws.append(['تعداد ستون‌ها', len(specs)])
        ws.append(['تاریخ تولید', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
        ws.append(['منبع', 'ExcelImporter.COLUMN_MAP'])
        ws.append([])
        ws.append(['ستون‌های ورودی'])
        for spec in specs:
            ws.append([spec['header'], spec['field']])
        ws.column_dimensions['A'].width = 28
        ws.column_dimensions['B'].width = 36

    @classmethod
    def generate_template_bytes(cls) -> Tuple[bytes, str, str]:
        wb, version = cls.build_workbook()
        buffer = BytesIO()
        wb.save(buffer)
        filename = f'قالب_ورودی_خرید_{version}.xlsx'
        return buffer.getvalue(), filename, version