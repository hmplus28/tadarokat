"""خروجی اکسل کامل همه رکوردهای خرید — قالب یکپارچه با ورودی/خروجی importer"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.excel_importer import ExcelImporter
from services.excel_template_service import FIELD_HINTS, ExcelTemplateService


PANEL_FIELD_HINTS = {
    '_deadline_deviated_label': 'بله/خیر — انحراف از مهلت درخواست (محاسبه پنل)',
    '_order_deadline_deviated_label': 'بله/خیر — انحراف از مهلت سفارش (محاسبه پنل)',
    '_payment_lead_deviated_label': 'بله/خیر — انحراف لید تایم پرداخت (محاسبه پنل)',
    '_current_status_label': 'وضعیت فعلی محاسبه‌شده در پنل — فقط خواندنی',
}


class ExcelFullExportService:
    """ساخت فایل اکسل خروجی کامل از تمام رکوردهای Purchase"""

    @classmethod
    def get_column_specs(cls) -> List[dict]:
        specs = []
        for header, field in ExcelImporter.get_export_columns():
            is_panel = field.startswith('_')
            specs.append({
                'header': header,
                'field': field,
                'required': field in ExcelImporter.IMPORT_REQUIRED_FIELDS,
                'decimal': field in ExcelImporter.DECIMAL_FIELDS,
                'workflow': field in ExcelImporter.WORKFLOW_FIELDS,
                'panel': is_panel,
                'aliases': ExcelImporter.get_import_aliases_for_field(field) if not is_panel else [],
                'hint': PANEL_FIELD_HINTS.get(field) or FIELD_HINTS.get(field, ''),
            })
        return specs

    @classmethod
    def _cell_export_value(cls, field: str, value) -> object:
        if field in ExcelImporter.ALWAYS_EXPORT_TEXT_FIELDS or field.startswith('_'):
            return value if value is not None else ''
        if value is None or value == '':
            return None
        return value

    @classmethod
    def build_workbook(cls) -> Tuple[Workbook, int]:
        from purchases.models import Purchase

        wb = Workbook()
        ws = wb.active
        ws.title = 'خروجی'

        specs = cls.get_column_specs()
        headers = [spec['header'] for spec in specs]
        fields = [spec['field'] for spec in specs]

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill('solid', fgColor='4F46E5')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        row_count = 0
        purchases = Purchase.objects.all().order_by('purchase_number', 'line_number', 'id')
        for purchase in purchases:
            row = []
            for field in fields:
                raw = ExcelImporter._field_export_value(purchase, field)
                row.append(cls._cell_export_value(field, raw))
            ws.append(row)
            row_count += 1

        for col_idx, header in enumerate(headers, start=1):
            width = max(len(header) + 2, 14)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 36)

        ws.freeze_panes = 'A2'
        cls._write_guide_sheet(wb, specs)
        cls._write_meta_sheet(wb, specs, row_count)
        return wb, row_count

    @classmethod
    def _write_guide_sheet(cls, wb: Workbook, specs: List[dict]) -> None:
        ws = wb.create_sheet('راهنما')
        ws.append(['ستون اکسل', 'فیلد سیستم', 'الزامی', 'نوع', 'نام‌های جایگزین', 'راهنما'])
        for spec in specs:
            if spec['panel']:
                type_label = 'پنل · فقط خروجی'
            else:
                type_label = 'عدد' if spec['decimal'] else 'متن/تاریخ'
                if spec['workflow']:
                    type_label += ' · گردش‌کار'
            aliases = '، '.join(spec['aliases']) if spec['aliases'] else '—'
            ws.append([
                spec['header'],
                spec['field'],
                'بله' if spec['required'] else 'خیر',
                type_label,
                aliases,
                spec['hint'] or '—',
            ])
        for col in range(1, 7):
            ws.column_dimensions[get_column_letter(col)].width = 22 if col < 6 else 40
        ws.freeze_panes = 'A2'

    @classmethod
    def _write_meta_sheet(cls, wb: Workbook, specs: List[dict], row_count: int) -> None:
        ws = wb.create_sheet('نسخه قالب')
        version = ExcelTemplateService.get_schema_version()
        ws.append(['کلید', 'مقدار'])
        ws.append(['نوع فایل', 'خروجی کامل دیتابیس'])
        ws.append(['نسخه اسکیما', version])
        ws.append(['تعداد ستون‌ها', len(specs)])
        ws.append(['تعداد ردیف‌ها', row_count])
        ws.append(['تاریخ تولید', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
        ws.append(['منبع', 'ExcelImporter.get_export_columns()'])
        ws.append([])
        ws.append(['ستون‌های خروجی'])
        for spec in specs:
            ws.append([spec['header'], spec['field']])
        ws.column_dimensions['A'].width = 28
        ws.column_dimensions['B'].width = 36

    @classmethod
    def generate_export_bytes(cls) -> Tuple[bytes, str, int]:
        wb, row_count = cls.build_workbook()
        buffer = BytesIO()
        wb.save(buffer)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'خروجی_کامل_خرید_{timestamp}.xlsx'
        return buffer.getvalue(), filename, row_count